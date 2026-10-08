from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# Create a presentation object
prs = Presentation()

# Define some colors
bg_color = RGBColor(33, 47, 61)  # Dark Blue/Grey
title_color = RGBColor(241, 196, 15) # Yellow/Gold
text_color = RGBColor(236, 240, 241) # Off-white

def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = bg_color

def add_title_slide(title_text, subtitle_text):
    slide_layout = prs.slide_layouts[0] # Title slide layout
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide)
    
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = title_text
    title.text_frame.paragraphs[0].font.color.rgb = title_color
    title.text_frame.paragraphs[0].font.bold = True
    
    subtitle.text = subtitle_text
    subtitle.text_frame.paragraphs[0].font.color.rgb = text_color

def add_bullet_slide(title_text, bullets):
    slide_layout = prs.slide_layouts[1] # Title and Content layout
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide)
    
    title = slide.shapes.title
    title.text = title_text
    title.text_frame.paragraphs[0].font.color.rgb = title_color
    title.text_frame.paragraphs[0].font.bold = True
    
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.text = bullets[0]
    tf.paragraphs[0].font.color.rgb = text_color
    
    for bullet in bullets[1:]:
        p = tf.add_paragraph()
        p.text = bullet
        p.font.color.rgb = text_color
        p.level = 0

# Slide 1: Title
add_title_slide("Avatarsi: AI-Powered Study Assistant", "Transforming Education with Intelligent Planning\nCollege Project Presentation")

# Slide 2: What is Avatarsi?
add_bullet_slide("What is Avatarsi?", [
    "A Java Spring Boot web application for education.",
    "Bridges the gap between scattered materials and structured learning.",
    "Acts as a smart companion that understands documents.",
    "Plans study schedules automatically."
])

# Slide 3: The Problem
add_bullet_slide("The Problem Statement", [
    "Information Overload: Massive PDFs take hours to read.",
    "Poor Time Management: Students struggle to plan study schedules.",
    "Resource Discovery: Finding quality tutorials is time-consuming.",
    "Information Retrieval: Searching inside PDFs is difficult."
])

# Slide 4: Key Features
add_bullet_slide("Core Capabilities", [
    "Automated Note Summarization: Upload PDFs for instant summaries.",
    "Dynamic Study Plan Generation: Topic & exam date based roadmap.",
    "Smart Document Search: Search concepts across all materials.",
    "Curated Video Recommendations: Suggests relevant YouTube tutorials."
])

# Slide 5: Value for Teaching
add_bullet_slide("Value for Teaching & Education", [
    "Curriculum Planning: Divide syllabus based on exam dates.",
    "Instant Handouts: Generate summaries from lecture notes.",
    "Flipped Classroom: Attach YouTube links for pre-class assignments.",
    "Student Support: Quickly locate concepts in course materials."
])

# Slide 6: Technical Architecture
add_bullet_slide("Technical Architecture", [
    "Backend Framework: Java Spring Boot.",
    "Document Processing: Apache PDFBox.",
    "AI & Logic Layer: Custom AIService handling text & dates.",
    "Frontend: HTML/CSS for a clean UI.",
    "Design Pattern: MVC (Model-View-Controller)."
])

# Slide 7: Conclusion
add_bullet_slide("Conclusion", [
    "Avatarsi is a comprehensive ecosystem for education.",
    "Saves time for teachers and provides structure for students.",
    "Makes learning targeted and efficient.",
    "\"Smart learning is about working with the right tools.\""
])

prs.save("Avatarsi_Presentation.pptx")
print("Presentation generated successfully at Avatarsi_Presentation.pptx")
